from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _excel_value(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        text = f"'{text}"
    return text[:32767]


def build_excel(title: str, headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "NEXORA"
    ws.freeze_panes = "A2"

    ws.append(headers)
    fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row in rows:
        ws.append([_excel_value(value) for value in row])

    for index, header in enumerate(headers, start=1):
        lengths = [len(str(header)), *(len(str(row[index - 1] or "")) for row in rows), 8]
        width = max(lengths)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 42)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_pdf(
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[list[object]],
    column_widths: list[float] | None = None,
) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
        author="NEXORA",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "NexoraTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=20,
        textColor=colors.HexColor("#0B3B36"),
        alignment=0,
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "NexoraSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=8,
    )
    cell_style = ParagraphStyle(
        "NexoraCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor("#1E293B"),
    )
    head_style = ParagraphStyle(
        "NexoraHead",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    data = [[Paragraph(escape(str(header)), head_style) for header in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(value if value is not None else "")), cell_style) for value in row])

    table = Table(data, colWidths=column_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    story = [
        Paragraph(escape(title), title_style),
        Paragraph(escape(subtitle), subtitle_style),
        Spacer(1, 3 * mm),
        table,
    ]
    doc.build(story)
    return buffer.getvalue()
